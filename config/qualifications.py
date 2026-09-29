"""资质档案 - 本地有真实 Excel 就加载，没有就用演示数据（公网部署用）"""

import os

from config.qualifications_demo import (
    DEMO_COMPANY,
    DEMO_QUALIFICATIONS,
    DEMO_SOFTWARE_COPYRIGHTS,
    DEMO_PATENTS,
)

# Excel 文件路径（本地真实数据，已被 .gitignore 忽略，不会上传公网）
_EXCEL_FILE = os.path.join(os.path.dirname(__file__), "产业互联网资质清单.xlsx")

def _load_qualifications_from_excel():
    """从 Excel 加载资质清单，失败返回 None"""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(_EXCEL_FILE, data_only=True)
        result = {"资质证书": [], "软件著作权": [], "专利": []}

        # Sheet1: 资质清单（第4行是表头）
        if "资质清单" in wb.sheetnames:
            ws = wb["资质清单"]
            for row in ws.iter_rows(min_row=5, values_only=True):
                if not row[0] or not isinstance(row[0], (int, float)):
                    continue
                result["资质证书"].append({
                    "序号": int(row[0]),
                    "资质名称": str(row[1] or ""),
                    "资质证书编号": str(row[2] or ""),
                    "资质有效期": str(row[3] or ""),
                    "发证机关": str(row[4] or ""),
                    "发证日期": str(row[5] or ""),
                })

        # Sheet2: 版权登记（第4行是表头）
        if "版权登记" in wb.sheetnames:
            ws = wb["版权登记"]
            for row in ws.iter_rows(min_row=5, values_only=True):
                if not row[0] or not isinstance(row[0], (int, float)):
                    continue
                result["软件著作权"].append({
                    "序号": int(row[0]),
                    "名称": str(row[1] or ""),
                    "类别": str(row[2] or ""),
                    "编号": str(row[3] or ""),
                    "有效期": str(row[4] or ""),
                })

        # Sheet3: 专利清单（第4行是表头）
        if "专利清单" in wb.sheetnames:
            ws = wb["专利清单"]
            for row in ws.iter_rows(min_row=5, values_only=True):
                if not row[0] or not isinstance(row[0], (int, float)):
                    continue
                result["专利"].append({
                    "序号": int(row[0]),
                    "名称": str(row[1] or ""),
                    "类型": str(row[2] or ""),
                    "专利号": str(row[3] or ""),
                })

        if result["资质证书"]:
            return result
    except Exception:
        pass
    return None


# 尝试加载本地真实数据；没有 Excel 就用演示数据
_REAL_DATA = _load_qualifications_from_excel()

if _REAL_DATA is not None:
    QUALIFICATIONS = _REAL_DATA["资质证书"]
    SOFTWARE_COPYRIGHTS = _REAL_DATA["软件著作权"]
    PATENTS = _REAL_DATA["专利"]
else:
    QUALIFICATIONS = DEMO_QUALIFICATIONS
    SOFTWARE_COPYRIGHTS = DEMO_SOFTWARE_COPYRIGHTS
    PATENTS = DEMO_PATENTS

# 投标主体信息：本地可在 qualifications_demo.py 或此处自行替换为真实信息
COMPANY = DEMO_COMPANY

# 资质名称列表（用于快速匹配）
QUALIFICATION_NAMES = [q["资质名称"] for q in QUALIFICATIONS]

# 在川标杆业绩
PROJECTS = [
    {"名称": "智慧蓉城城市大脑二期（成都高新区）", "金额": "3200万", "类型": "智慧城市/城市大脑"},
    {"名称": "四川省大数据技术服务中心数据中台", "金额": "5800万", "类型": "大数据平台"},
    {"名称": "成都市某区智慧政务一体化平台", "金额": "1800万", "类型": "智慧政务"},
    {"名称": "四川省某厅局政务云资源服务", "金额": "4200万/年", "类型": "政务云"},
    {"名称": "成都天府新区城市运营管理平台", "金额": "2600万", "类型": "城市运营"},
    {"名称": "四川省应急管理厅智慧应急平台", "金额": "2800万", "类型": "智慧应急"},
    {"名称": "成都市武侯区智慧社区综合治理平台", "金额": "1500万", "类型": "智慧社区"},
    {"名称": "四川省医保局智能审核系统", "金额": "2200万", "类型": "智慧医疗"},
    {"名称": "内江市大数据中心数据治理项目", "金额": "1900万", "类型": "大数据平台"},
    {"名称": "泸州市智慧文旅一卡通平台", "金额": "1200万", "类型": "智慧文旅"},
]

# 人员配置
PERSONNEL = {
    "项目经理": "PMP/高项 50人+",
    "网络安全专家": "CISSP/CISP 30人+",
    "大数据工程师": "100人+",
}

# 核心产品
PRODUCTS = [
    "联通云（新沃云）",
    "资治政务大数据平台",
    "城市大脑",
    "AI能力平台（AI数字民警、法智惠企）",
]


def format_qualifications_md():
    """把资质业绩格式化成 Markdown 文本，用于 Prompt 注入"""
    lines = [
        f"【我方主体：{COMPANY['主体全称']}】",
        f"【统一社会信用代码：{COMPANY['统一社会信用代码']}】",
        "",
        f"## 核心资质（共{len(QUALIFICATIONS)}项）",
    ]
    for q in QUALIFICATIONS:
        lines.append(
            f"- {q['资质名称']} | 编号:{q['资质证书编号']} | "
            f"有效期:{q['资质有效期']} | 发证:{q['发证机关']}"
        )

    if SOFTWARE_COPYRIGHTS:
        lines.append(f"\n## 软件著作权（共{len(SOFTWARE_COPYRIGHTS)}项，列举前10）")
        for sc in SOFTWARE_COPYRIGHTS[:10]:
            lines.append(f"- {sc['名称']}（{sc['编号']}）")

    if PATENTS:
        lines.append(f"\n## 专利（共{len(PATENTS)}项，列举前5）")
        for p in PATENTS[:5]:
            lines.append(f"- {p['名称']}（{p['类型']}，{p['专利号']}）")

    lines += [
        "",
        "## 在川标杆业绩",
    ]
    for i, p in enumerate(PROJECTS, 1):
        lines.append(f"{i}. {p['名称']}，{p['金额']}，类型：{p['类型']}")

    lines += [
        "",
        "## 人员配置",
        f"- 项目经理：{PERSONNEL['项目经理']}",
        f"- 网络安全专家：{PERSONNEL['网络安全专家']}",
        f"- 大数据工程师：{PERSONNEL['大数据工程师']}",
        "",
        "## 核心产品",
    ]
    for prod in PRODUCTS:
        lines.append(f"- {prod}")
    return "\n".join(lines)
