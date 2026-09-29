"""文档解析模块 - 支持 PDF 和 Word

特性：
- 100MB 大小限制
- magic number 文件头真校验（不只看扩展名）
- 扫描件 PDF 检测（文字层缺失率过高时给出 OCR 提示）
"""

import os
import tempfile
from dataclasses import dataclass


FILE_SIZE_LIMIT = 100 * 1024 * 1024  # 100MB

# 判定扫描件的阈值：超过 50% 的页面提取不到文字，即认为是扫描件
SCANNED_EMPTY_PAGE_RATIO = 0.5
SCANNED_MIN_CHARS_PER_PAGE = 10


@dataclass
class ParseResult:
    """解析结果"""
    text: str
    is_scanned: bool = False   # 是否疑似扫描件
    page_count: int = 0        # PDF 页数（其他格式为 0）
    file_ext: str = ""


def parse_pdf(file_path: str) -> tuple[str, int, bool]:
    """解析 PDF 文件，返回 (文本, 页数, 是否疑似扫描件)"""
    import pdfplumber
    text_parts = []
    total_pages = 0
    empty_pages = 0

    with pdfplumber.open(file_path) as pdf:
        total_pages = len(pdf.pages)
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            page_text = page_text.strip()
            if len(page_text) < SCANNED_MIN_CHARS_PER_PAGE:
                empty_pages += 1
            if page_text:
                text_parts.append(page_text)

    is_scanned = (
        total_pages > 0
        and (empty_pages / total_pages) >= SCANNED_EMPTY_PAGE_RATIO
    )
    return "\n\n".join(text_parts), total_pages, is_scanned


def parse_docx(file_path: str) -> str:
    """解析 Word 文件（段落 + 表格）"""
    from docx import Document
    doc = Document(file_path)
    text_parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            text_parts.append(para.text)
    # 表格内容
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            text_parts.append(" | ".join(cells))
    return "\n".join(text_parts)


def parse_txt(file_path: str) -> str:
    """解析纯文本文件（容忍编码问题）"""
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def _validate_file(uploaded_file, file_ext: str, raw: bytes) -> None:
    """大小 + 文件头 magic number 校验"""
    if len(raw) > FILE_SIZE_LIMIT:
        size_mb = len(raw) / 1024 / 1024
        raise ValueError(f"文件过大：{size_mb:.1f}MB，超过 100MB 上限，请拆分或压缩后上传")

    if file_ext == ".pdf" and not raw.startswith(b"%PDF"):
        raise ValueError("文件扩展名为 .pdf，但文件头不是有效 PDF（%PDF），请确认文件未损坏")

    if file_ext in (".docx",) and not raw.startswith(b"PK\x03\x04"):
        raise ValueError("文件扩展名为 .docx，但文件头不是有效 Word（ZIP/OOXML），请确认文件未损坏")

    if file_ext == ".txt":
        try:
            raw.decode("utf-8")
        except UnicodeDecodeError:
            # 尝试常见中文编码
            for enc in ("gb18030", "gbk", "gb2312"):
                try:
                    raw.decode(enc)
                    return
                except UnicodeDecodeError:
                    continue
            raise ValueError("TXT 文件编码无法识别，请另存为 UTF-8 编码后再上传")


def parse_uploaded_file(uploaded_file) -> ParseResult:
    """
    解析 Streamlit 上传的文件对象
    返回 ParseResult（文本 + 扫描件标记 + 页数）
    """
    file_ext = os.path.splitext(uploaded_file.name)[1].lower()
    raw = uploaded_file.getvalue()

    _validate_file(uploaded_file, file_ext, raw)

    # 写入临时文件供 pdfplumber / python-docx 使用
    with tempfile.NamedTemporaryFile(delete=False, suffix=file_ext) as tmp:
        tmp.write(raw)
        tmp_path = tmp.name

    try:
        if file_ext == ".pdf":
            text, page_count, is_scanned = parse_pdf(tmp_path)
            result = ParseResult(
                text=text,
                is_scanned=is_scanned,
                page_count=page_count,
                file_ext=file_ext,
            )
        elif file_ext == ".docx":
            result = ParseResult(text=parse_docx(tmp_path), file_ext=file_ext)
        elif file_ext == ".doc":
            raise ValueError(
                "旧版 .doc（Word 97-2003）格式暂不支持，请用 Word 另存为 .docx 后上传"
            )
        elif file_ext == ".txt":
            # UTF-8 优先，失败回退 gb18030
            try:
                text = raw.decode("utf-8")
            except UnicodeDecodeError:
                text = raw.decode("gb18030", errors="ignore")
            result = ParseResult(text=text, file_ext=file_ext)
        else:
            raise ValueError(f"不支持的文件格式：{file_ext}，请上传 PDF / DOCX / TXT")
    finally:
        try:
            os.unlink(tmp_path)
        except OSError:
            pass

    if not result.text.strip():
        if file_ext == ".pdf":
            result.is_scanned = True
            raise ValueError(
                "未从 PDF 中提取到任何文字，这通常是扫描件/图片版招标文件。"
                "请先使用 OCR 工具（如联通内网 OCR、WPS 会员 OCR、Adobe OCR）识别文字后再上传。"
            )
        raise ValueError("文件内容为空，请检查后重新上传")

    return result
